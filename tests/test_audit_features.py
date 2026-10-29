"""Demo data fixtures only. End-to-end source-file handling in temporary SQLite."""
import base64,hashlib,json
from io import BytesIO
from datetime import date
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from model.api.urban import create_app

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setenv('NAGARSEVA_OFFICER_KEY','audit-test-key')
    with TestClient(create_app(tmp_path/'audit.sqlite3',tmp_path/'no-source')) as c:yield c

HEADERS={'X-Officer-Key':'audit-test-key'}
def source(raw,name='demo-data.csv'):
    data=raw.encode() if isinstance(raw,str) else raw
    return {'filename':name,'contentBase64':base64.b64encode(data).decode(),'permissionAccepted':True}
def mapping(raw='id,name\nDEMO-IN-1,Demo data settlement\n',**changes):return {**source(raw),'category':'identifiers','mapping':{'id':'id','name':'name'},'crsConfirmed':False,'settlementBatchId':None,**changes}
def body(payload):return {**payload,'confirm':True,'expectedSha256':hashlib.sha256(base64.b64decode(payload['contentBase64'])).hexdigest(),'sourceName':'Demo data test source','publisher':'Demo data test suite','sourceDate':'2026-10-01','licence':'Demo data - test use only','sourceUrl':'https://example.test/demo-data','coverage':'Demo data fictional city','dataClass':'demo'}
def save(client,payload):return client.post('/urban/integration/imports',headers=HEADERS,json=body(payload))
def xlsx(rows):
    book=Workbook();sheet=book.active;sheet.title='Demo data'
    for row in rows:sheet.append(row)
    data=BytesIO();book.save(data);return data.getvalue()


def test_csv_template_is_real_and_contains_no_invented_records(client):
    result=client.get('/urban/integration/templates/services')
    assert result.status_code==200 and result.headers['content-type'].startswith('text/csv')
    assert result.text.splitlines()==['id,settlementId,measurementDate,waterAccessPercent,sanitationAccessPercent,drainageCoveragePercent,wasteCollectionPercent,electricityAccessPercent,roadAccessPercent']


def test_xlsx_valid_inspection_mapping_preview_and_preserved_bytes(client):
    raw=xlsx([['code','name'],['DEMO-XLSX-1','Demo data worksheet record']]);payload={**mapping(),**source(raw,'demo-data.xlsx'),'mapping':{'id':'code','name':'name'}}
    inspected=client.post('/urban/integration/inspect',json=payload_file(payload)).json()
    assert inspected['format']=='XLSX' and inspected['recordCount']==1 and inspected['sheet']=='Demo data' and inspected['columns']==['code','name']
    validated=client.post('/urban/integration/validate',json=payload).json()
    assert validated['valid'] and validated['preview'][0]['id']=='DEMO-XLSX-1' and validated['completenessPercent']>0
    created=save(client,payload);assert created.status_code==200,created.text
    receipt=created.json();assert receipt['is_synthetic'] is True and receipt['provenance']['licence']=='Demo data - test use only'
    assert client.get('/urban/integration/imports/'+receipt['id']+'/original',headers=HEADERS).content==raw


def payload_file(payload):return {key:payload[key] for key in ['filename','contentBase64','permissionAccepted']}


@pytest.mark.parametrize('rows,expected',[
    ([['id','name'],['DEMO-1','=1+1']],'audit.error.formula'),
    ([['id','id'],['DEMO-1','Demo data']],'ir.error.columns'),
    ([['id','name']],'ir.error.empty'),
    ([['id','name'],['DEMO-1','Demo data'],['DEMO-1','Demo data duplicate']],'ir.error.duplicate'),
])
def test_bad_workbooks_are_clear_and_never_saved(client,rows,expected):
    payload={**mapping(),**source(xlsx(rows),'demo-data.xlsx')}
    result=client.post('/urban/integration/validate',json=payload)
    assert result.status_code==422 and expected in result.text or result.status_code==200 and any(row['code']==expected for row in result.json()['errors'])
    assert sum(row['recordCount'] for row in client.get('/urban/integration/readiness').json()['categories'])==0


@pytest.mark.parametrize('field',['sourceName','publisher','sourceDate','licence','sourceUrl'])
def test_all_provenance_fields_are_required(client,field):
    data=body(mapping());del data[field]
    assert client.post('/urban/integration/imports',json=data,headers=HEADERS).status_code==422
    assert client.get('/urban/integration/imports',headers=HEADERS).json()['items']==[]


def test_csv_wrong_types_row_errors_duplicate_missing_empty_and_large(client):
    batch=save(client,mapping()).json()['id']
    payload={**source('id,settlement,count,date\nDEMO-H-1,DEMO-IN-1,not-a-number,2026-10-01\n'),'category':'housing','mapping':{'id':'id','settlementId':'settlement','householdCount':'count','assessmentDate':'date'},'settlementBatchId':batch,'crsConfirmed':False}
    result=client.post('/urban/integration/validate',json=payload,headers=HEADERS).json()
    assert any(error['code']=='ir.error.number' and error['row']==2 and error['field']=='householdCount' for error in result['errors'])
    for raw,expected in [('name\nDemo data\n','ir.error.mapping'),('id,name\nDUP,Demo data\nDUP,Demo data\n','ir.error.duplicate'),('id,name\n','ir.error.empty')]:
        result=client.post('/urban/integration/validate',json=mapping(raw)).json();assert not result['valid'] and any(error['code']==expected for error in result['errors'])
    for raw in [b'',b'x'*(5*1024*1024+1)]:
        response=client.post('/urban/integration/inspect',json=source(raw));assert response.status_code==422 and 'ir.error.size' in response.text
    assert len(client.get('/urban/integration/imports',headers=HEADERS).json()['items'])==1


def test_rollback_is_confirmed_audited_dependency_safe_and_excludes_links(client):
    parent=save(client,mapping()).json()['id']
    child_payload={**source('id,settlement,date,coverage\nDEMO-M-1,DEMO-IN-1,2026-10-01,75\n'),'category':'services','mapping':{'id':'id','settlementId':'settlement','measurementDate':'date','waterAccessPercent':'coverage'},'settlementBatchId':parent,'crsConfirmed':False}
    child=save(client,child_payload).json()['id']
    endpoint='/urban/integration/imports/'
    def rollback(identifier,confirm=True):return client.post(endpoint+identifier+'/rollback',json={'confirm':confirm,'reason':'Demo data rollback test'},headers=HEADERS)
    assert rollback(parent).status_code==409
    assert rollback(child,False).status_code==422
    assert rollback(child).status_code==200
    result=rollback(parent);assert result.status_code==200 and result.json()['history'][-1]['action']=='rollback'
    assert sum(row['importedBatches'] for row in client.get('/urban/integration/readiness').json()['categories'])==0
    assert client.get(endpoint+parent+'/original',headers=HEADERS).status_code==200
    assert client.post('/urban/integration/validate',json=child_payload,headers=HEADERS).status_code==422
    assert len(client.get('/urban/demo/workspace').json()['settlements'])==100


def test_no_official_classification_is_inferred_from_an_upload(client):
    payload=body(mapping());payload['dataClass']='official'
    assert client.post('/urban/integration/imports',json=payload,headers=HEADERS).status_code==422


def test_synthetic_flags_planted_cases_and_measured_large_workload(client):
    before=client.get('/urban/demo/workspace').json()
    for key in ['settlements','wards','issues','projects','milestones','visits','activity','dataSources','anomalyFixtures']:
        assert all(row['is_synthetic'] is True for row in before[key])
    assert len(before['anomalyFixtures'])==3
    result=client.get('/urban/demo/benchmark?count=10000');assert result.status_code==200,result.text
    measured=result.json();assert measured['recordCount']==10000 and measured['totalMs']>0 and measured['is_synthetic'] and not measured['persisted']
    assert client.get('/urban/demo/workspace').json()==before


def test_xlsx_percent_cells_keep_displayed_units_and_sheet_choice(client):
    book=Workbook();sheet=book.active;sheet.title='Demo data';sheet.append(['id','coverage']);sheet.append(['DEMO-PCT-1',.75]);sheet['B2'].number_format='0.0%'
    second=book.create_sheet('Demo data second');second.append(['id','name']);second.append(['DEMO-SHEET-2','Demo data second worksheet'])
    stream=BytesIO();book.save(stream);payload=source(stream.getvalue(),'Demo-data.xlsx')
    result=client.post('/urban/integration/inspect',json=payload).json()
    assert result['excelPercentColumns']==['coverage'] and len(result['sheets'])==2
    from model.api.file_formats import read_xlsx
    records,_,_,_=read_xlsx(stream.getvalue());assert records[0]['coverage']==75
    result=client.post('/urban/integration/inspect',json={**payload,'sheetName':'Demo data second'}).json()
    assert result['sheet']=='Demo data second' and result['columns']==['id','name']
    assert client.post('/urban/integration/inspect',json={**payload,'sheetName':'Not a worksheet'}).status_code==422
