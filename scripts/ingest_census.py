"""Download official Census metadata and advertised assets, retaining 2011 geography.

No estimates are allocated from state/town/ward totals to settlement clusters.
Excel tables are preserved cell-for-cell; numerical observation adapters require
reviewed headers and explicit geography keys, not positional guesses.
"""
from __future__ import annotations
import argparse
import hashlib
from io import BytesIO
from pathlib import Path
import ssl
import subprocess
from urllib.parse import urljoin, urlparse
from bs4 import BeautifulSoup
import xlrd
from openpyxl import load_workbook
from data_common import Downloader, DownloadError, META, PROCESSED, ROOT, inventory, utcnow, write_json
from census_tables import parse_slum_pca, pca_maharashtra, parse_housing, parse_water, parse_water_electricity_latrine

CATALOGUES = {
    "6190": "Census 2011 slum PCA", "11436": "Census 2011 slum house condition",
    "11448": "Census 2011 water electricity latrine", "11443": "Census 2011 slum drinking water",
}
INTERMEDIATE_URL = "https://repository.emsign.com/certs/emSignSSLCAG1.crt"
INTERMEDIATE_SHA256 = "47b2efbc3670e7db4b41f22c51fc02ee84fb2dbf3082a49f2c2688122e9210a1"


def census_tls_context(refresh=False):
    """Complete the server's missing intermediate, retaining OS-root validation.

    The official public CA intermediate is fetched using ordinary HTTPS first,
    pinned, and independently verified to an installed root with OpenSSL. No new
    root CA is installed and certificate/hostname checking remain mandatory.
    """
    downloader = Downloader("census", refresh)
    try:
        raw, request = downloader.fetch(INTERMEDIATE_URL, "census-emsign-intermediate", suffix=".crt")
    finally:
        downloader.close()
    if hashlib.sha256(raw).hexdigest() != INTERMEDIATE_SHA256:
        raise ValueError("Census intermediate changed; review public CA certificate before use")
    verification = subprocess.run(["openssl", "verify", str(ROOT / request["path"])], capture_output=True, text=True, timeout=15, check=True)
    context = ssl.create_default_context()
    context.load_verify_locations(cadata=ssl.DER_cert_to_PEM_cert(raw))
    report = {"source_url": INTERMEDIATE_URL, "sha256": INTERMEDIATE_SHA256, "request": request,
              "verification": verification.stdout.strip(), "method": "OS-root verified intermediate completes the server chain; hostname validation and CERT_REQUIRED retained", "check_hostname": context.check_hostname, "verify_mode": "CERT_REQUIRED"}
    write_json(META / "census_tls.json", report)
    return context, report


def historical_unit(state, town, name, source):
    resolution = "state" if town == "000000" else "town"
    identifier = f"census2011-{state}-{town}"
    return {"id": identifier, "source_id": f"{state}:{town}", "name": name, "resolution": resolution,
            "reference_year": 2011, "source": source, "state_code": state, "town_code": town,
            "geometry": None, "parent_id": None if resolution == "state" else f"census2011-{state}-000000",
            "geographic_key_system": "Census of India 2011 State Code / Town Code", "synthetic": False,
            "limitations": "Original 2011 geography; no current-boundary or settlement crosswalk; source labels including (Part) retained"}


def observation(unit, field, value, indicator_type, source, request, catalogue, **extra):
    return {"id": f"census2011-{catalogue}-{unit['source_id'].replace(':', '-')}-{field}", "settlement_id": None,
            "geographic_unit_id": unit["id"], "indicator_type": indicator_type, "value": value,
            "unit": "persons" if indicator_type == "population" else "households", "reference_year": 2011,
            "observation_date": None, "geographic_resolution": unit["resolution"], "source": source,
            "verification_status": "historical_published", "inspector_reference": "published-source",
            "method": "Exact published Census cell; original geography and reference year retained; not a current service deficit index",
            "evidence_reference": None, "source_field": field, "source_label": field, "catalogue_id": catalogue,
            "retrieved_at": request["retrieved_at"], "raw_path": request["path"], "raw_sha256": request["sha256"],
            "synthetic": False, **extra}


def parse_workbook(raw, suffix):
    if suffix == ".xls":
        book = xlrd.open_workbook(file_contents=raw)
        sheets = [{"name": s.name, "rows": [s.row_values(i) for i in range(s.nrows)]} for s in book.sheets()]
    else:
        book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        sheets = [{"name": s.title, "rows": [[v.isoformat() if hasattr(v, "isoformat") else v for v in row] for row in s.iter_rows(values_only=True)]} for s in book]
    if not sheets or not any(s["rows"] for s in sheets):
        raise ValueError("Empty Census workbook")
    if sum(len(s["rows"]) for s in sheets) > 100_000:
        raise ValueError("Workbook exceeds pilot row limit")
    return sheets


def ingest(refresh=False):
    try:
        context, tls = census_tls_context(refresh)
    except (DownloadError, ValueError, subprocess.SubprocessError) as exc:
        # Normal TLS remains safe if certificate acquisition is unavailable.
        context, tls = True, {"error": str(exc), "method": "Unmodified default HTTPS verification; no certificate bypass"}
    downloader = Downloader("census", refresh, verify=context)
    report = {"attempted_at": utcnow(), "tables": [], "errors": [], "reference_year": 2011}
    report["tls"] = tls
    units, observations, tables = {}, [], []
    for code, name in CATALOGUES.items():
        url = f"https://censusindia.gov.in/nada/index.php/catalog/{code}"
        entry = {"catalogue_id": code, "name": name, "source_url": url, "assets": []}
        try:
            raw, request = downloader.fetch(url, "catalogue-" + code, suffix=".html")
            soup = BeautifulSoup(raw, "html.parser")
            text = soup.get_text(" ", strip=True)
            if "2011" not in text or "Download" not in text or "Census" not in text:
                raise ValueError("Census catalogue content not present; refusing to treat an error page as metadata")
            links = list(dict.fromkeys(urljoin(url, a["href"]) for a in soup.select("a[href]") if urlparse(a["href"]).path.lower().endswith((".xls", ".xlsx", ".pdf"))))
            entry.update(request=request, advertised_assets=links, metadata_text=text)
            for i, asset_url in enumerate(links[:8]):
                suffix = Path(urlparse(asset_url).path).suffix.lower()
                try:
                    content, info = downloader.fetch(asset_url, f"table-{code}-{i}", suffix=suffix)
                    if suffix == ".pdf":
                        if not content.startswith(b"%PDF-"):
                            raise ValueError("Not a PDF")
                        output = subprocess.run(["pdftotext", "-layout", str(ROOT / info["path"]), "-"], capture_output=True, text=True, check=True, timeout=30).stdout
                        if len(output.strip()) < 50:
                            raise ValueError("No extractable Census table text; OCR/manual review required")
                        parsed = {"reference_year": 2011, "raw_text": output, "numeric_observations_activated": False, "request": info}
                        count = None
                        validation = "PDF text parsed; table values/geography not automatically validated or activated"
                        adapter = {"11436": parse_housing, "11443": parse_water, "11448": parse_water_electricity_latrine}.get(code)
                        if adapter:
                            cells = adapter(output)
                            unit = historical_unit("27", "000000", "MAHARASHTRA", "https://censusindia.gov.in/nada/index.php/catalog/6190")
                            units[unit["id"]] = unit
                            for cell in cells:
                                kind = "household_amenities" if code == "11448" else cell["indicator_type"]
                                observations.append(observation(unit, cell["source_field"], cell["value"], kind, asset_url, info, code,
                                    source_label=cell["source_label"], denominator=cell["denominator"], denominator_reference_year=2011,
                                    universe=cell["universe"], table_phase="Houselisting and Housing Census", missing_value_semantics="Printed dash remains null; not zero"))
                            parsed.update(numeric_observations_activated=True, geography="Maharashtra state27 only", cells=cells)
                            count = len(cells)
                            validation = "Reviewed Maharashtra state27 cell adapter; header, column count and arithmetic subtotals validated; original PDF retained"
                    else:
                        sheets = parse_workbook(content, suffix)
                        parsed = {"reference_year": 2011, "sheets": sheets, "request": info, "numeric_observations_activated": False}
                        count = sum(len(s["rows"]) for s in sheets)
                        validation = "Workbook cells parsed; row count includes headers; geography adapter requires review"
                        if code == "6190":
                            records = parse_slum_pca(sheets)
                            subset = pca_maharashtra(records)
                            for row in subset:
                                unit = historical_unit(row["State Code"], row["Town Code"], row["Area Name"], url)
                                units[unit["id"]] = unit
                                for field, kind in (("No_HH", "households"), ("TOT_P", "population")):
                                    observations.append(observation(unit, field, row[field], kind, asset_url, info, code,
                                        source_label="Slum households" if kind == "households" else "Slum population",
                                        universe="Population and households in slums reported by Census 2011 PCA", table_phase="Population Enumeration / Primary Census Abstract"))
                            parsed.update(records=records, numeric_observations_activated=True, geographic_key_fields=["State Code", "Town Code"])
                            write_json(PROCESSED / "census_slum_pca_maharashtra.json", {"reference_year": 2011, "request": info, "columns": list(records[0]), "records": subset})
                            count = len(records)
                            validation = f"{len(records)} original geography rows / 88 columns schema; unique zero-padded codes, nonnegative integers, population identities and Maharashtra state/town totals validated"
                    tables.append({"catalogue_id": code, "request": info, "record_count": count, "validation": validation})
                    write_json(PROCESSED / f"census_{code}_{i}.json", parsed)
                    entry["assets"].append({"request": info, "record_count": count, "validation": validation})
                    inventory(dataset_name=name + " asset " + str(i), publisher="ORGI Census of India", source_url=asset_url, retrieved_at=info["retrieved_at"],
                              observation_date="2011", geography="Source table geography retained; not allocated to clusters", resolution="Catalogue/table-specific", format=suffix[1:],
                              record_count=count, fields=list(parsed.get("records", [{}])[0]) if parsed.get("records") else [cell["source_field"] for cell in parsed.get("cells", [])], licence="Official public Census tables; no additional third-party licence asserted", access_status="downloaded", validation_status=validation,
                              intended_use="Historical context only", limitations="2011 is not current; no cluster allocation; no inferred service reliability")
                except (DownloadError, ValueError, subprocess.SubprocessError) as exc:
                    entry.setdefault("errors", []).append({"url": asset_url, "error": str(exc)})
            inventory(dataset_name=name + " metadata", publisher="ORGI Census of India", source_url=url, retrieved_at=request["retrieved_at"], observation_date="2011",
                      format="HTML", access_status="metadata_only", validation_status="Catalogue text parsed; " + str(len(links)) + " advertised files", intended_use="Discover historical table and geography", limitations="Catalogue creation/modification dates are not observation dates")
        except (DownloadError, ValueError) as exc:
            entry["error"] = str(exc)
            report["errors"].append({"catalogue_id": code, "error": str(exc)})
            inventory(dataset_name=name + " metadata", publisher="ORGI Census of India", source_url=url, retrieved_at=utcnow(), observation_date="2011",
                      access_status="failed", validation_status="No local table parsed", intended_use="Historical context only", limitations=str(exc))
        report["tables"].append(entry)
    for key, url in (("population-finder", "https://censusindia.gov.in/census.website/data/population-finder"), ("dataful-maharashtra", "https://dataful.in/datasets/272/")):
        try:
            raw, info = downloader.fetch(url, key, suffix=".html")
            text = BeautifulSoup(raw, "html.parser").get_text(" ", strip=True)
            entry = {"request": info, "content_summary": text[:2000]}
            report[key] = entry
            inventory(dataset_name=key, publisher="ORGI" if key == "population-finder" else "Dataful / Factly (secondary convenience source)", source_url=url,
                      retrieved_at=info["retrieved_at"], observation_date="2011", format="HTML", access_status="metadata_only", validation_status="HTML parsed; no structured table obtained",
                      intended_use="Discovery only", limitations="No downloadable table verified; dynamic/registration access not bypassed; use official table for analysis")
        except DownloadError as exc:
            inventory(dataset_name=key, publisher="ORGI" if key == "population-finder" else "Dataful / Factly", source_url=url, retrieved_at=utcnow(), access_status="failed", validation_status="not parsed", limitations=str(exc))
    observation_data = {"geographic_units": list(units.values()), "indicators": observations,
                        "metadata": {"reference_year": 2011, "synthetic": False, "geographic_unit_count": len(units), "indicator_count": len(observations),
                                     "tables": tables, "limitations": ["Historical 2011 aggregates only; no settlement allocation or current service inference.", "PCA and Houselisting household universes differ: Maharashtra2499948 versus2449530. Never share denominators.", "Missing PDF cells remain null; original table labels and raw checksums retained."]}}
    write_json(PROCESSED / "census_observations.json", observation_data)
    report["processed"] = observation_data["metadata"]
    write_json(META / "census_ingestion.json", report)
    downloader.close()
    print("Census: ", sum(len(t["assets"]) for t in report["tables"]), "assets;", len(report["errors"]), "catalogue errors")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--refresh", action="store_true")
    ingest(parser.parse_args().refresh)
