"""Generate a concise deck describing implemented behavior without invented impact figures."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
NAVY = RGBColor(9, 42, 67)
TEAL = RGBColor(21, 157, 154)
INK = RGBColor(23, 60, 84)
MUTED = RGBColor(103, 127, 131)


def text(slide, value, x, y, width, height, size=20, colour=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    for index, line in enumerate(value.split('\n')):
        paragraph = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
        paragraph.text = line
        paragraph.space_after = Pt(13)
        paragraph.font.name = 'Aptos'
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = colour
        paragraph.font.bold = bold
    return box


def build():
    deck = Presentation(); deck.slide_width = Inches(13.333); deck.slide_height = Inches(7.5)
    slides = [
        ('NagarSeva', 'Plan housing improvements.\nVerify residents receive the promised services.', ['Housing planning + household affordability', 'Accountable action + resident verification', 'Working software; a field pilot is the next step']),
        ('The problem', 'A home needs affordable living costs and functioning services.', ['Compare housing options alongside monthly household costs.', 'Connect a service problem to a responsible department and deadline.', 'Record evidence and keep the resident involved in resolution.']),
        ('One connected project', 'A shared record across the whole workflow.', ['Priority map → project brief → financial and household scenario', 'Resident report → officer assignment → dated repair measurement', 'Resident confirmation or dispute → updated service record']),
        ('Affordability on both sides', 'Developer viability and household burden stay visible.', ['Project: area, rehab assumptions, market rates, construction and sensitivity.', 'Household: rent/EMI, maintenance, utilities, transport and income loss.', 'Temporary rent and moving costs are separate; all assumptions are editable.']),
        ('Working service assurance', 'A repair is submitted for resident review.', ['SQLite records survive restarts; each resident receives a private receipt.', 'Officers review category/urgency, assign work and record after-repair evidence.', 'Confirmation accepts the reading. Dispute reopens the report and excludes it.']),
        ('Credible allocation', 'Show the constraints and an actual baseline.', ['Distinct city/ward coverage, budget and vulnerable-share constraints are checked.', 'Equal-ward budget allocation is independently solved with its caps exposed.', 'Impact is a risk-weighted planning proxy; no realized savings are claimed.']),
        ('Evidence and limits', 'Sources and sample sizes make the demonstration reviewable.', ['Public context: OpenStreetMap, dated Open-Meteo forecasts and official guidance.', 'Planning labels remain synthetic; forecasts do not validate settlement risk.', 'Consented field observations and user trials form separate pilot evidence.']),
        ('Pilot request', 'One approved project. A resident partner. A municipal reviewer.', ['Validate task completion, response, resident-confirmed resolution and measured services.', 'Record source-backed project commitments and evidence limitations.', 'PMAY-U 2.0 source: pmaymis.gov.in/pmaymis2_2024/PmayFAQ.aspx · reviewed 2026-10-11']),
    ]
    for number, (eyebrow, title, bullets) in enumerate(slides, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        text(slide, f'{number:02} / {eyebrow.upper()}', .65, .4, 12, .4, 12, TEAL, True)
        text(slide, title, .65, 1.05, 11.9, 1.55, 30, NAVY, True)
        screenshot = {4: 'affordability-desktop.png', 5: 'assurance-viewport.png'}.get(number)
        asset = ROOT / 'presentation/assets' / screenshot if screenshot else None
        has_screenshot = asset is not None and asset.exists()
        for index, bullet in enumerate(bullets):
            text(slide, f'{index + 1:02}', .7, 3.0 + index * 1.03, .5, .6, 15, TEAL, True)
            text(slide, bullet, 1.4, 3.0 + index * 1.03, 4.6 if has_screenshot else 11.0, .95, 17 if has_screenshot else 19)
        if has_screenshot:
            picture = slide.shapes.add_picture(str(asset), Inches(6.6), Inches(2.65), width=Inches(6.0))
            if picture.height > Inches(3.9):
                picture.width = int(picture.width * Inches(3.9) / picture.height)
                picture.height = Inches(3.9)
            picture.left = Inches(6.6) + (Inches(6.0) - picture.width) // 2
            text(slide, 'Captured from the tested app · synthetic demo records', 6.6, 6.7, 6.0, .25, 10, MUTED)
        text(slide, 'NagarSeva · Synthetic planning baseline · Pilot evidence is source-labelled', .65, 7.08, 12, .25, 9, MUTED)
    output = ROOT / 'presentation/NagarSeva_Pitch.pptx'
    deck.save(output)
    print(f'Created {output.name}: {len(slides)} slides')


if __name__ == '__main__':
    build()
