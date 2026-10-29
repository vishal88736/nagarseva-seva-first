"""Bounded, read-only XLSX ingestion. No formula evaluation or external links."""
from datetime import date,datetime,time
import re
from io import BytesIO
from zipfile import ZipFile,BadZipFile
from openpyxl import load_workbook
from fastapi import HTTPException


def read_xlsx(raw,sheet_name=None):
    def reject(code):raise HTTPException(422,code)
    try:
        with ZipFile(BytesIO(raw)) as archive:
            files=archive.infolist()
            if len(files)>1000 or sum(info.file_size for info in files)>20*1024*1024:reject('audit.error.archive')
            for item in files:
                if item.filename.endswith('.xml'):
                    data=archive.read(item)
                    if b'<!DOCTYPE' in data or b'<!ENTITY' in data:reject('audit.error.workbook')
            if any('vbaproject' in item.filename.lower() or 'externallinks/' in item.filename.lower() or item.flag_bits&1 for item in files):reject('audit.error.workbook')
        book=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        try:
            names=[sheet.title for sheet in book.worksheets if sheet.sheet_state=='visible']
            if not names or sheet_name and sheet_name not in names:reject('audit.error.sheet')
            selected=sheet_name or names[0];sheet=book[selected]
            if (sheet.max_row or 0)>5001 or (sheet.max_column or 0)>100:reject('ir.error.limit')
            iterator=sheet.iter_rows(max_row=min(sheet.max_row or 1,5002),max_col=min(sheet.max_column or 1,101))
            headers=next(iterator,())
            columns=[cell.value for cell in headers]
            while columns and columns[-1] is None:columns.pop()
            if not columns or any(not isinstance(key,str) or not key.strip() for key in columns) or len(set(columns))!=len(columns):reject('ir.error.columns')
            records=[];errors=[];percent_columns=set()
            for source_row,cells in enumerate(iterator,2):
                if all(cell.value is None for cell in cells):continue
                if len(records)>=5000:reject('ir.error.limit')
                row={}
                for key,cell in zip(columns,cells):
                    value=cell.value
                    if cell.data_type in ('f','e'):
                        errors.append({'code':'audit.error.formula','row':source_row,'field':key,'detail':None});value=None
                    elif isinstance(value,datetime):value=value.date().isoformat() if value.time()==time() else value.isoformat()
                    elif isinstance(value,date):value=value.isoformat()
                    if isinstance(value,(int,float)) and not isinstance(value,bool) and '%' in re.sub(r'"[^"]*"|\\.', '',cell.number_format):
                        value*=100;percent_columns.add(key)
                    row[key]=value
                # Retain actual worksheet row numbers for the validation report.
                row['@sourceRow']=source_row;records.append(row)
            return records,columns,errors,{'sheets':names,'sheet':selected,'excelPercentColumns':sorted(percent_columns)}
        finally:book.close()
    except HTTPException:raise
    except Exception as exc:raise HTTPException(422,'audit.error.workbook') from exc
